# -*- coding: utf-8 -*-
"""
Extractor for the official conference-attendance form (HR-SD-S-F13).

Unlike the free-form memo (text_pipeline), this document is a FIXED FORM:
key-value pairs separated by tabs, a fixed 6-line expense section, a
traveler list, and a 'กำหนดการเดินทาง' schedule.

Conventions (see docs/HOW_TO_ADD_NEW_FORM.md):
  - Main line: 'ข้าพเจ้า <tab> ชื่อ <tab> ตำแหน่ง <tab> ... <tab> เรื่อง <tab> ชื่องาน
                <tab> ระหว่างวันที่ X ถึงวันที่ Y ... จังหวัด <tab> จ. <tab> จัดโดย <tab> หน่วยงาน
                <tab> โดยมีจำนวนผู้เดินทางทั้งหมด N คน'
  - Travelers: header 'รายชื่อผู้ร่วมเดินทาง' then rows:
                'ลำดับ <tab> คำนำหน้า <tab> ชื่อ <tab> นามสกุล <tab> ตำแหน่ง <tab> หน่วยงาน'
  - Expenses:  6 fixed lines 'ค่าลงทะเบียน ... เป็นเงิน X บาท' etc.
  - Breakdown: section 'ประมาณการค่าใช้จ่าย' (second one) with plaintext rows.
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List

import docx

from ..normalizer.thai_utils import parse_thai_date
from .text_pipeline import _parse_plaintext_breakdown, _parse_schedule_activities

# 6 fixed expense categories in the official form (order matters)
EXPENSE_CATEGORIES = [
    "ค่าลงทะเบียน",
    "ค่าเบี้ยเลี้ยง",
    "ค่าที่พัก",
    "ค่าพาหนะเดินทาง",
    "เงินชดเชยพาหนะส่วนตัว",
    "ค่าธรรมเนียม/ค่าใช้จ่ายอื่นๆ",
]
CATEGORY_KEYS = ["registration", "per_diem", "accommodation", "vehicle", "vehicle_compensation", "other"]

# Event type keywords -> form select values
EVENT_TYPE_MAP = [
    ("ประชุม", "conference"),
    ("สัมมนา", "seminar"),
    ("ฝึกอบรม", "training"),
    ("อบรม", "training"),
    ("workshop", "workshop"),
    ("ศึกษาดูงาน", "site-visit"),
    ("ดูงาน", "site-visit"),
    ("ไปราชการ", "official-travel"),
    ("อื่นๆ", "other"),
]

VEHICLE_KEYWORDS = ("พาหนะ", "ตั๋ว", "เครื่องบิน", "รถไฟ", "เดินทาง", "รถโดยสาร", "แท็กซี่", "รถตู้")


def _split_tokens(line: str) -> List[str]:
    return [t.strip() for t in line.split("\t") if t.strip()]


def _derive_event_type(title: str) -> str:
    t = (title or "").lower()
    for kw, val in EVENT_TYPE_MAP:
        if kw in t:
            return val
    return "other"


def _parse_main_line(main_line: str, extracted: Dict[str, Any]) -> None:
    """Parse the tab-separated 'ข้าพเจ้า ...' line into structured fields."""
    tokens = _split_tokens(main_line)
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        nxt = tokens[i + 1] if i + 1 < len(tokens) else ""
        if tok == "ข้าพเจ้า" and nxt:
            extracted["requester_name"] = nxt
        elif "ตำแหน่ง" in tok and nxt and not extracted.get("requester_position"):
            extracted["requester_position"] = nxt
        elif "ภาควิชา/กอง/ส่วน/ศูนย์/งาน" in tok and nxt:
            extracted["department"] = nxt
        elif "คณะ/สำนัก" in tok and nxt:
            extracted["faculty"] = nxt
        elif "เรื่อง" in tok and nxt and not extracted.get("event_title"):
            extracted["event_title"] = nxt
            extracted["project_title"] = nxt  # ชื่องาน = ชื่อโครงการสำหรับฟอร์มนี้
        elif "ระหว่างวันที่" in tok:
            # e.g. 'ระหว่างวันที่ 3 ธันวาคม 2568 ถึงวันที่ 3 ธันวาคม 2568 เป็นเวลา 1 วัน สถานที่จัดประชุม/อบรม/สัมมนา จังหวัด'
            dm = re.search(
                r"ระหว่างวันที่\s*(\d{1,2}\s+[^\s\d]+\s+\d{4})\s*ถึงวันที่\s*(\d{1,2}\s+[^\s\d]+\s+\d{4})",
                tok,
            )
            if dm:
                d1, d2 = parse_thai_date(dm.group(1)), parse_thai_date(dm.group(2))
                if d1:
                    extracted["start_date_iso"] = d1.strftime("%Y-%m-%d")
                if d2:
                    extracted["end_date_iso"] = d2.strftime("%Y-%m-%d")
                extracted["schedule_text"] = f"{dm.group(1)} ถึง {dm.group(2)}"
            if "จังหวัด" in tok:
                extracted["province_name"] = nxt or extracted.get("province_name", "")
        elif "จัดโดย" in tok and nxt:
            extracted["organizer"] = nxt
        elif "จำนวนผู้เดินทางทั้งหมด" in tok:
            cm = re.search(r"(\d+)", tok)
            if cm:
                extracted["traveler_count"] = cm.group(1)
        i += 1
    if extracted.get("province_name"):
        loc = extracted.get("location_name") or ""
        extracted["location_province"] = f"ณ {loc} จ.{extracted['province_name']}".strip()


def _parse_expense_section(full_text: List[str], extracted: Dict[str, Any]) -> None:
    """Parse the 6 fixed expense lines '... เป็นเงิน X บาท' + รวม."""
    cats: Dict[str, float] = {}
    total = 0.0
    in_section = False
    for p in full_text:
        if "ประมาณการค่าใช้จ่ายประกอบด้วย" in p:
            in_section = True
            continue
        if in_section:
            if p.startswith("รวม") or p.startswith("จึงเรียนมา") or p.startswith("ลงชื่อ"):
                break
            m = re.search(r"^(ค่าลงทะเบียน|ค่าเบี้ยเลี้ยง|ค่าที่พัก|ค่าพาหนะเดินทาง|เงินชดเชยพาหนะส่วนตัว|ค่าธรรมเนียม/ค่าใช้จ่ายอื่นๆ).*?เป็นเงิน\s*([\d,]*)\s*บาท", p)
            if m:
                key = CATEGORY_KEYS[EXPENSE_CATEGORIES.index(m.group(1))]
                amt = float(m.group(2).replace(",", "")) if m.group(2) else 0.0
                cats[key] = amt
    extracted["expense_categories"] = {k: cats.get(k, 0.0) for k in CATEGORY_KEYS}

    # รวม line (may be after the section or within)
    for p in full_text:
        m = re.search(r"^รวม.*?เป็นเงิน\s*([\d,]+)\s*บาท", p)
        if m:
            total = float(m.group(1).replace(",", ""))
            break
    if total > 0:
        extracted["budget_amount"] = str(int(total))
        extracted["budget_text"] = ""
    else:
        extracted["budget_amount"] = str(int(sum(cats.values()))) if sum(cats.values()) else ""


def _parse_travelers(paragraphs: List[str], extracted: Dict[str, Any]) -> None:
    """Parse 'รายชื่อผู้ร่วมเดินทาง' section rows."""
    travelers: List[Dict[str, str]] = []
    in_section = False
    for p in paragraphs:
        if p.strip() == "รายชื่อผู้ร่วมเดินทาง":
            in_section = True
            continue
        if in_section:
            if p.startswith("กำหนดการ") or p.startswith("ประมาณการ") or p.startswith("ความเห็น"):
                break
            parts = _split_tokens(p)
            # expected: [ลำดับ, คำนำหน้า, ชื่อ, นามสกุล, ตำแหน่ง, หน่วยงาน]
            if len(parts) >= 4:
                travelers.append({
                    "no": parts[0],
                    "title": parts[1],
                    "first_name": parts[2],
                    "last_name": parts[3],
                    "position": parts[4] if len(parts) > 4 else "",
                    "department": parts[5] if len(parts) > 5 else "",
                })
    if travelers:
        extracted["travelers"] = travelers
        extracted["traveler_count"] = str(len(travelers))


def _parse_breakdown(full_text: List[str], extracted: Dict[str, Any]) -> None:
    """Parse the SECOND 'ประมาณการค่าใช้จ่าย' section (plaintext rows)."""
    idxs = [i for i, p in enumerate(full_text) if p.strip() == "ประมาณการค่าใช้จ่าย"]
    if not idxs:
        return
    start = idxs[-1]
    section = full_text[start + 1:]
    # stop at รวม line
    cut = next((i for i, p in enumerate(section) if p.startswith("รวม")), len(section))
    rows = _parse_plaintext_breakdown(section[:cut])
    if rows:
        extracted["breakdown"] = rows


def _parse_vehicle_details(extracted: Dict[str, Any]) -> None:
    """textarea-83 content: vehicle-related breakdown rows, one per line."""
    lines = []
    for row in extracted.get("breakdown") or []:
        item = str(row.get("รายการ", ""))
        if any(kw in item for kw in VEHICLE_KEYWORDS):
            detail = str(row.get("รายละเอียด", ""))
            line = f"{item} ({detail})" if detail and detail != "-" else item
            lines.append(line)
    extracted["vehicle_details"] = "\n".join(lines)


def _parse_approval_table(doc, extracted: Dict[str, Any]) -> None:
    """Approval chain from the native table (ผู้บังคับบัญชา levels)."""
    chain = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            name = ""
            pos = ""
            for c in cells:
                nm = re.search(r"\(?\s*(นาย|นาง|นางสาว|ดร\.|ผศ\.|รศ\.|ศ\.)\s*([^\)]+)\)?", c)
                if nm and not name:
                    name = f"{nm.group(1)}{nm.group(2).strip()}".strip("() ")
                elif c and len(c) < 60 and c not in ("อนุมัติ", "ไม่อนุมัติ", "ความเห็นเพิ่มเติม", "ลงนาม", "ตำแหน่ง"):
                    pos = c
            if name:
                chain.append({"level": cells[0] if cells else "", "name": name, "position": pos})
    if chain:
        extracted["approval_chain"] = chain
        if len(chain) >= 1:
            extracted["approver_left_name"] = chain[0]["name"]
            extracted["approver_left_pos"] = chain[0]["position"]


def extract_conference(file_bytes: bytes) -> Dict[str, Any]:
    """Parse the conference-attendance form (.docx) into the normalized dict."""
    doc = docx.Document(io.BytesIO(file_bytes))

    paragraphs: List[str] = []
    full_text: List[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            paragraphs.append(t)
            full_text.append(t)

    extracted: Dict[str, Any] = {
        "doc_type_hint": "conference_attendance",
        "agency_name": "",
        "contact_phone": "",
        "doc_number": "",
        "doc_number_tail": "",
        "doc_date": "",
        "doc_date_iso": "",
        "budget_year": "",
        "recipient_title": "",
        "closing_text": "",
        "requester_name": "",
        "requester_position": "",
        "department": "",
        "faculty": "",
        "project_title": "",
        "event_title": "",
        "organizer": "",
        "traveler_count": "1",
        "travelers": [],
        "location_name": "",
        "province_name": "",
        "location_province": "",
        "schedule_text": "",
        "start_date_iso": "",
        "end_date_iso": "",
        "budget_amount": "",
        "budget_text": "",
        "expense_categories": {},
        "vehicle_details": "",
        "breakdown": [],
        "schedule_activities": [],
        "approval_chain": [],
        "action_verb": "เข้าร่วม",
    }

    # วันที่ / เรื่อง / เรียน (only the FIRST วันที่ line is the doc date —
    # schedule day lines like 'วันที่ 3 ธันวาคม 2568' must not overwrite it)
    for p in paragraphs:
        m = re.match(r"^วันที่\s*[:：]?\s*(\d{1,2}\s+[^\s\d]+\s+\d{4})", p)
        if m and not extracted.get("doc_date"):
            extracted["doc_date"] = m.group(1).strip()
            d = parse_thai_date(extracted["doc_date"])
            if d:
                extracted["doc_date_iso"] = d.strftime("%Y-%m-%d")
                extracted["budget_year"] = str((d.year + 543) % 100)
            continue
        if p.startswith("เรื่อง") and not extracted.get("event_title"):
            extracted["project_title"] = re.sub(r"^เรื่อง\s*[:：]?\s*", "", p).strip()
            continue
        if p.startswith("เรียน") and not extracted.get("recipient_title"):
            extracted["recipient_title"] = re.sub(r"^เรียน\s*[:：]?\s*", "", p).strip()
            continue
        if p.startswith("ข้าพเจ้า"):
            _parse_main_line(p, extracted)
            continue

    # 6 fixed expense categories + รวม
    _parse_expense_section(full_text, extracted)

    # Travelers
    _parse_travelers(paragraphs, extracted)

    # Breakdown (second ประมาณการค่าใช้จ่าย section)
    _parse_breakdown(full_text, extracted)

    # Vehicle details for textarea
    _parse_vehicle_details(extracted)

    # Event type (for the form select)
    extracted["event_type"] = _derive_event_type(extracted.get("event_title") or extracted.get("project_title", ""))

    # Schedule + approval table
    extracted["schedule_activities"] = _parse_schedule_activities(full_text, extracted)
    _parse_approval_table(doc, extracted)

    return extracted

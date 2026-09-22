# -*- coding: utf-8 -*-
"""DOCX ingestion via python-docx."""
from __future__ import annotations

import io
from typing import Any, Dict, List

import docx

from .text_pipeline import extract_from_text


def _iter_tables(tables) -> List[List[List[str]]]:
    """Convert python-docx tables to a plain list structure."""
    out: List[List[List[str]]] = []
    for table in tables:
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        out.append(rows)
    return out


def extract_docx(file_bytes: bytes) -> Dict[str, Any]:
    """Parse a .docx memo and return the extracted field dictionary."""
    doc = docx.Document(io.BytesIO(file_bytes))

    paragraphs: List[str] = []
    full_text: List[str] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)
            full_text.append(txt)

    tables = _iter_tables(doc.tables)
    return extract_from_text(paragraphs, full_text, tables)

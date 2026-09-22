# -*- coding: utf-8 -*-
"""
Excel (.xlsx) ingestion — reads the label→value template sheets.

Template convention (see docs/HOW_TO_ADD_NEW_FORM.md):
  Sheet 'ฟอร์ม': column A = field label, column B = value.
  Section header 'รายชื่อผู้ร่วมเดินทาง' switches the parser to traveler rows
  (ลำดับ/คำนำหน้า/ชื่อ/นามสกุล/ตำแหน่ง/หน่วยงาน).
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List

from openpyxl import load_workbook

from ..normalizer.thai_utils import parse_thai_date

# Label (exact, lowercased) -> normalized key
LABEL_MAP = {
    "ชื่องาน/กิจกรรม": "event_title",
    "วันที่เริ่มต้น": "start_date",
    "วันที่สิ้นสุด": "end_date",
    "สถานที่": "location_name",
    "จังหวัด": "province_name",
    "หน่วยงานผู้จัด": "organizer",
    "ค่าลงทะเบียน (บาท)": "registration",
    "ค่าเบี้ยเลี้ยง (บาท)": "per_diem",
    "ค่าที่พัก (บาท)": "accommodation",
    "ค่าพาหนะเดินทาง (บาท)": "vehicle",
    "เงินชดเชยพาหนะส่วนตัว (บาท)": "vehicle_compensation",
    "ค่าธรรมเนียม/ค่าใช้จ่ายอื่นๆ (บาท)": "other",
    "รายละเอียดค่าพาหนะเดินทาง": "vehicle_details",
    "ชื่อ-นามสกุล": "requester_name",
    "ตำแหน่ง": "requester_position",
    "ภาควิชา/กอง/ส่วน/ศูนย์/งาน": "department",
    "คณะ/สำนัก": "faculty",
    "วันที่หนังสือ": "doc_date",
}

TRAVELER_HEADER = "รายชื่อผู้ร่วมเดินทาง"


def _norm(s) -> str:
    s = str(s or "").replace("\n", " ").strip().lower()
    s = re.sub(r"^\d+\.\s*", "", s)  # strip "1. " / "2. " numbering prefixes
    return s


def _match_label(label: str):
    """Match a sheet label against LABEL_MAP (exact, then contains)."""
    n = _norm(label)
    if n in LABEL_MAP:
        return LABEL_MAP[n]
    for key, val in LABEL_MAP.items():
        if key in n or n in key:
            return val
    return None


def extract_excel(file_bytes: bytes) -> Dict[str, Any]:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb["ฟอร์ม"] if "ฟอร์ม" in wb.sheetnames else wb.active

    values: Dict[str, Any] = {}
    travelers: List[Dict[str, str]] = []
    in_travelers = False

    for row in ws.iter_rows(min_row=1, values_only=True):
        label = str(row[0] or "").strip()
        value = row[1] if len(row) > 1 else None
        if not label:
            continue

        if label == TRAVELER_HEADER:
            in_travelers = True
            continue
        if in_travelers:
            # traveler row: [no, title, first, last, position, department]
            cells = [str(c or "").strip() for c in row[:6]]
            # Only count a row as a traveler when a NAME is present
            # (template pre-fills the ลำดับ numbers 1..5 which are not travelers)
            if (cells[2] or cells[3]) and any(cells):
                travelers.append({
                    "no": cells[0], "title": cells[1], "first_name": cells[2],
                    "last_name": cells[3], "position": cells[4] if len(cells) > 4 else "",
                    "department": cells[5] if len(cells) > 5 else "",
                })
            continue

        key = _match_label(label)
        if key:
            values[key] = value

    # Normalize dates -> ISO
    for dk, iso_k in (("start_date", "start_date_iso"), ("end_date", "end_date_iso")):
        if values.get(dk):
            d = parse_thai_date(str(values[dk]))
            values[iso_k] = d.strftime("%Y-%m-%d") if d else str(values[dk])

    if values.get("doc_date"):
        d = parse_thai_date(str(values["doc_date"]))
        if d:
            values["doc_date_iso"] = d.strftime("%Y-%m-%d")

    # expense_categories + total
    cats = {k: float(str(values.get(k) or 0).replace(",", "")) for k in
            ("registration", "per_diem", "accommodation", "vehicle", "vehicle_compensation", "other")}
    values["expense_categories"] = cats
    total = sum(cats.values())
    if total > 0:
        values["budget_amount"] = str(int(total))

    # project_title = event_title
    values["project_title"] = values.get("event_title", "")
    if values.get("province_name"):
        loc = values.get("location_name", "")
        values["location_province"] = f"ณ {loc} จ.{values['province_name']}".strip()
    values["travelers"] = travelers
    values["traveler_count"] = str(len(travelers)) if travelers else "1"

    return values

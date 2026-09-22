# -*- coding: utf-8 -*-
"""Excel export (openpyxl) — pre-filled workbook for the office template."""
from __future__ import annotations

import io
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

THAI_FONT = "Sarabun"
HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(name=THAI_FONT, size=12, bold=True, color="FFFFFF")
BODY_FONT = Font(name=THAI_FONT, size=12)
TITLE_FONT = Font(name=THAI_FONT, size=16, bold=True)


def _style_header(ws, row: int, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def generate_excel_bytes(extracted: Dict[str, Any], rollup: Dict[str, Any]) -> bytes:
    """Build the workbook and return its bytes."""
    wb = Workbook()

    # ---- Sheet 1: ข้อมูลโครงการ ----
    ws1 = wb.active
    ws1.title = "ข้อมูลโครงการ"
    ws1["A1"] = "ข้อมูลโครงการ (RSC Auto-Extract)"
    ws1["A1"].font = TITLE_FONT

    meta_rows = [
        ("ส่วนงาน", extracted.get("agency_name", "")),
        ("เลขที่หนังสือ", extracted.get("doc_number", "")),
        ("วันที่หนังสือ", extracted.get("doc_date", "")),
        ("เรื่อง/ชื่อโครงการ", extracted.get("project_title", "")),
        ("ผู้ขอ", extracted.get("requester_name", "")),
        ("ตำแหน่ง", extracted.get("requester_position", "")),
        ("สถานที่ดำเนินการ", extracted.get("location_province", "")),
        ("ช่วงเวลา", extracted.get("schedule_text", "")),
        ("เริ่มต้น (ISO)", extracted.get("start_date_iso", "")),
        ("สิ้นสุด (ISO)", extracted.get("end_date_iso", "")),
        ("กลุ่มเป้าหมาย", f"{extracted.get('target_group_name', '')} {extracted.get('target_group_quantity', '')} {extracted.get('target_group_unit', '')}".strip()),
        ("วงเงินรวม (บาท)", extracted.get("budget_amount", "")),
        ("วงเงินตัวอักษร", extracted.get("budget_text", "")),
    ]
    for i, (k, v) in enumerate(meta_rows, start=3):
        ws1.cell(row=i, column=1, value=k).font = Font(name=THAI_FONT, size=12, bold=True)
        ws1.cell(row=i, column=2, value=str(v)).font = BODY_FONT
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 80

    # ---- Sheet 2: ประมาณการค่าใช้จ่าย ----
    ws2 = wb.create_sheet("ประมาณการค่าใช้จ่าย")
    ws2["A1"] = "ประมาณการค่าใช้จ่าย"
    ws2["A1"].font = TITLE_FONT
    headers = ["ลำดับ", "รายการ", "รายละเอียด", "จำนวนเงิน (บาท)", "หมวดหมู่"]
    for col, h in enumerate(headers, start=1):
        ws2.cell(row=3, column=col, value=h)
    _style_header(ws2, 3, len(headers))

    rows = rollup.get("rows") or []
    for i, row in enumerate(rows, start=1):
        ws2.cell(row=3 + i, column=1, value=i)
        ws2.cell(row=3 + i, column=2, value=row.get("รายการ", ""))
        ws2.cell(row=3 + i, column=3, value=row.get("รายละเอียด", ""))
        ws2.cell(row=3 + i, column=4, value=row.get("จำนวนเงิน (บาท)", 0))
        ws2.cell(row=3 + i, column=5, value=row.get("หมวดหมู่", ""))
        for col in range(1, 6):
            ws2.cell(row=3 + i, column=col).font = BODY_FONT
            ws2.cell(row=3 + i, column=4).number_format = "#,##0.00"

    total_row = 4 + len(rows)
    ws2.cell(row=total_row, column=3, value="รวมเป็นเงินทั้งสิ้น").font = Font(name=THAI_FONT, size=12, bold=True)
    ws2.cell(row=total_row, column=4, value=rollup.get("total_amount", 0)).font = Font(name=THAI_FONT, size=12, bold=True)
    ws2.cell(row=total_row, column=4).number_format = "#,##0.00"
    ws2.cell(row=total_row + 1, column=3, value=f"({num_to_thai_text(rollup.get('total_amount', 0))})").font = BODY_FONT

    for col, width in zip(range(1, 6), [8, 42, 42, 18, 22]):
        ws2.column_dimensions[get_column_letter(col)].width = width
    ws2.freeze_panes = "A4"

    # ---- Sheet 3: กำหนดการ ----
    ws3 = wb.create_sheet("กำหนดการ")
    ws3["A1"] = "กำหนดการ"
    ws3["A1"].font = TITLE_FONT
    headers3 = ["วันที่", "สถานที่", "เวลา", "กิจกรรม"]
    for col, h in enumerate(headers3, start=1):
        ws3.cell(row=3, column=col, value=h)
    _style_header(ws3, 3, len(headers3))

    r = 4
    for day in extracted.get("schedule_activities") or []:
        items = day.get("items") or []
        if not items:
            ws3.cell(row=r, column=1, value=day.get("date_title", ""))
            ws3.cell(row=r, column=2, value=day.get("location", ""))
            for col in range(1, 5):
                ws3.cell(row=r, column=col).font = BODY_FONT
            r += 1
            continue
        for act in items:
            ws3.cell(row=r, column=1, value=day.get("date_title", ""))
            ws3.cell(row=r, column=2, value=day.get("location", ""))
            ws3.cell(row=r, column=3, value=act.get("time", ""))
            ws3.cell(row=r, column=4, value=act.get("activity", ""))
            for col in range(1, 5):
                ws3.cell(row=r, column=col).font = BODY_FONT
                ws3.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=True)
            r += 1

    for col, width in zip(range(1, 5), [26, 34, 18, 60]):
        ws3.column_dimensions[get_column_letter(col)].width = width
    ws3.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def num_to_thai_text(amount: float) -> str:
    """Small wrapper to avoid a circular import from normalizer."""
    from ..normalizer.thai_utils import num_to_thai_baht

    return num_to_thai_baht(amount)


# ---------------------------------------------------------------------------
# Fill-in Excel template (Quick Form / Excel mode)
# ---------------------------------------------------------------------------
def generate_fill_template_bytes(fields) -> bytes:
    """Build a label→value .xlsx template from a profile's field schema.

    fields: list of {"key", "label", "type"} (from /api/v1/forms).
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ฟอร์ม"
    ws["A1"] = "ฟิลด์"
    ws["B1"] = "ค่าที่กรอก"
    ws["A1"].font = Font(name=THAI_FONT, size=12, bold=True)
    ws["B1"].font = Font(name=THAI_FONT, size=12, bold=True)

    r = 2
    for f in fields:
        label = str(f.get("label", ""))
        key = str(f.get("key", ""))
        # ใช้ label สั้น (ไม่มีเลขหมวด) เป็นคอลัมน์ A เพื่อให้ excel_extractor อ่านกลับได้
        ws.cell(row=r, column=1, value=label).font = BODY_FONT
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        t = f.get("type", "text")
        if t == "number":
            ws.cell(row=r, column=2, value=0).number_format = "#,##0.00"
        elif t == "textarea":
            ws.cell(row=r, column=2, value="").alignment = Alignment(wrap_text=True, vertical="top")
        elif t == "date":
            ws.cell(row=r, column=2, value="2026-01-01")
        r += 1

    # Traveler section (convention)
    ws.cell(row=r, column=1, value="รายชื่อผู้ร่วมเดินทาง").font = Font(name=THAI_FONT, size=12, bold=True)
    r += 1
    headers = ["ลำดับ", "คำนำหน้า", "ชื่อ", "นามสกุล", "ตำแหน่ง", "หน่วยงาน"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=r, column=c, value=h).font = Font(name=THAI_FONT, size=12, bold=True)
    for i in range(1, 6):
        ws.cell(row=r + i, column=1, value=i)
        for c in range(1, 7):
            ws.cell(row=r + i, column=c).font = BODY_FONT

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

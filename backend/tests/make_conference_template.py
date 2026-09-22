# -*- coding: utf-8 -*-
"""
Generate the blank conference-attendance DOCX template (HR-SD-S-F13 convention).

Run:
    python tests/make_conference_template.py [output.docx]

The layout follows the extractor conventions (see docs/HOW_TO_ADD_NEW_FORM.md):
  - Main line: tab-separated key/value pairs (ข้าพเจ้า/ตำแหน่ง/ภาควิชา/คณะ/เรื่อง/ระหว่างวันที่/จังหวัด/จัดโดย/จำนวนผู้เดินทาง)
  - 6 fixed expense lines '... เป็นเงิน [ ] บาท'
  - 'รายชื่อผู้ร่วมเดินทาง' section
  - 'กำหนดการเดินทาง' + 'ประมาณการค่าใช้จ่าย' sections
"""
from __future__ import annotations

import os
import sys

import docx
from docx.shared import Pt

TAB = "\t"


def build(path: str) -> None:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "TH Sarabun New"
    style.font.size = Pt(14)

    p = doc.add_paragraph()
    r = p.add_run("แบบขออนุมัติเข้าร่วมประชุม / อบรม / สัมมนา ในประเทศ")
    r.bold = True

    doc.add_paragraph("มหาวิทยาลัยเทคโนโลยีพระจอมเกล้าธนบุรี")

    doc.add_paragraph(f"วันที่ {TAB * 2}[วัน] [เดือน] [พ.ศ.]")
    doc.add_paragraph(f"เรื่อง {TAB}ขออนุมัติเข้าร่วมประชุม/อบรม/สัมมนาในประเทศ")
    doc.add_paragraph(f"เรียน{TAB}[ผู้บังคับบัญชา/ผู้อำนวยการ]")
    doc.add_paragraph()

    # Main line (tab-separated)
    main = (
        f"ข้าพเจ้า {TAB}[ชื่อ-นามสกุล]{TAB}ตำแหน่ง {TAB}[ตำแหน่ง]{TAB}"
        f"ภาควิชา/กอง/ส่วน/ศูนย์/งาน{TAB}[ภาควิชา/กอง/ส่วน/ศูนย์/งาน]{TAB}"
        f"คณะ/สำนัก{TAB}[คณะ/สำนัก]{TAB}"
        f"มีความประสงค์จะเข้าร่วมประชุม/อบรม/สัมมนา เรื่อง {TAB}[ชื่องาน/กิจกรรม]{TAB}"
        f"ระหว่างวันที่ [วัน] [เดือน] [พ.ศ.] ถึงวันที่ [วัน] [เดือน] [พ.ศ.] เป็นเวลา [N] วัน "
        f"สถานที่จัดประชุม/อบรม/สัมมนา จังหวัด{TAB}[จังหวัด]{TAB}"
        f"จัดโดย{TAB}[หน่วยงานผู้จัด]{TAB}"
        f"โดยมีจำนวนผู้เดินทางทั้งหมด [N] คน"
    )
    doc.add_paragraph(main)
    doc.add_paragraph()

    # Fixed expense section
    doc.add_paragraph("ประมาณการค่าใช้จ่ายประกอบด้วย")
    for cat in ("ค่าลงทะเบียน", "ค่าเบี้ยเลี้ยง", "ค่าที่พัก", "ค่าพาหนะเดินทาง",
                "เงินชดเชยพาหนะส่วนตัว", "ค่าธรรมเนียม/ค่าใช้จ่ายอื่นๆ"):
        doc.add_paragraph(f"{cat}{TAB * 4}เป็นเงิน {TAB}[จำนวน]{TAB}บาท")
    doc.add_paragraph(f"รวมค่าใช้จ่ายทั้งหมด {TAB}(ขอถัวเฉลี่ยทุกรายการ){TAB}เป็นเงิน {TAB}[รวม]{TAB}บาท")
    doc.add_paragraph()

    # Traveler list
    doc.add_paragraph("รายชื่อผู้ร่วมเดินทาง")
    doc.add_paragraph(f"ลำดับ{TAB}คำนำหน้า{TAB}ชื่อ{TAB}นามสกุล{TAB}ตำแหน่ง{TAB}หน่วยงาน")
    for i in range(1, 6):
        doc.add_paragraph(f"{i}{TAB}{TAB}{TAB}{TAB}{TAB}{TAB}")
    doc.add_paragraph()

    # Schedule
    doc.add_paragraph("กำหนดการเดินทาง")
    doc.add_paragraph("วันที่ [วัน] [เดือน] [พ.ศ.]")
    for _ in range(4):
        doc.add_paragraph(f"[เวลา]{TAB}[กิจกรรม]")
    doc.add_paragraph()

    # Breakdown
    doc.add_paragraph("ประมาณการค่าใช้จ่าย")
    for _ in range(4):
        doc.add_paragraph(f"[รายการ]{TAB * 6}[จำนวน]{TAB}บาท")
    doc.add_paragraph(f"รวม{TAB * 8}[รวม]{TAB}บาท")

    doc.save(path)
    print("saved:", path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates",
        "แบบขออนุมัติเข้าร่วมประชุม_blank.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    build(out)

# -*- coding: utf-8 -*-
"""
Generate the blank RSC memo (บันทึกข้อความ) DOCX template used as the
canonical hardcoded form for template-anchored extraction (see
backend/templates/rsc_memo.json).

Variable slots are marked with {{key}} placeholders; everything else is fixed
boilerplate that the extractor cuts out of filled documents.

Run:
    python tests/make_memo_template.py [output.docx]

Layout follows the RSC memo skeleton observed in the filled sample documents
(ส่วนงาน / ที่ อว. / วันที่ / เรื่อง / เรียน / ตามที่… / ดังนั้น… / ในการนี้…
/ จึงเรียนมา… / ลายเซ็น / ประมาณการค่าใช้จ่าย / กำหนดการ).
"""
from __future__ import annotations

import os
import sys

import docx
from docx.shared import Inches, Pt

TAB = "\t"
DOTS = "........................"


def _repo_root() -> str:
    return os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def build(path: str) -> None:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "TH Sarabun New"
    style.font.size = Pt(14)

    # Header: logo + บันทึกข้อความ
    logo = os.path.join(_repo_root(), "kmutt_logo.png")
    if os.path.exists(logo):
        try:
            p = doc.add_paragraph()
            p.alignment = 1  # center
            p.add_run().add_picture(logo, width=Inches(1.0))
        except Exception:  # noqa: BLE001 — logo is cosmetic
            pass

    p = doc.add_paragraph()
    p.alignment = 1
    r = p.add_run("บันทึกข้อความ")
    r.bold = True
    r.font.size = Pt(22)

    doc.add_paragraph(f"ส่วนงาน {TAB}{{{{agency_name}}}} {TAB}โทร {{{{contact_phone}}}}")
    doc.add_paragraph(f"ที่ อว. {{{{doc_number}}}} {TAB}วันที่ {{{{doc_date}}}}")
    doc.add_paragraph(f"เรื่อง {TAB}{{{{project_title}}}}")
    doc.add_paragraph(DOTS * 6)
    doc.add_paragraph()
    doc.add_paragraph(f"เรียน {TAB}{{{{recipient_title}}}}")
    doc.add_paragraph(f"เอกสารแนบ {TAB}1. ประมาณการค่าใช้จ่าย 2. รายละเอียดกำหนดการ")
    doc.add_paragraph()

    # เนื้อหา
    doc.add_paragraph(
        "ตามที่ศูนย์ส่งเสริมและสนับสนุนมูลนิธิโครงการหลวง และโครงการตามพระราชดำริ "
        "สถาบันพัฒนาและฝึกอบรมโรงงานต้นแบบ มหาวิทยาลัยเทคโนโลยีพระจอมเกล้าธนบุรี "
        f"ได้ดำเนินงาน {{{{project_context}}}}"
    )
    doc.add_paragraph(
        f"ดังนั้น {{{{project_objective}}}} ข้าพเจ้า {{{{requester_name}}}} "
        f"ตำแหน่ง {{{{requester_position}}}} ขออนุมัติดำเนินงาน {{{{action_details}}}} "
        f"ณ {{{{location_name}}}} จังหวัด {{{{province_name}}}} "
        f"ระหว่างวันที่ {{{{schedule_text}}}} โดยมีกลุ่มเป้าหมายประกอบด้วย "
        f"{{{{target_group_name}}}} จำนวน {{{{target_group_quantity}}}} {{{{target_group_unit}}}}"
    )
    doc.add_paragraph(
        "ในการนี้ข้าพเจ้าจึงใคร่ขออนุมัติดำเนินงานตามวัน เวลา และสถานที่ดังกล่าว "
        f"พร้อมทั้งขออนุมัติค่าใช้จ่าย จำนวน {{{{budget_amount}}}} บาท ({{{{budget_text}}}}) "
        f"ขอถัวเฉลี่ยทุกรายการ จากงบประมาณ {{{{budget_source}}}}"
    )
    doc.add_paragraph(f"จึงเรียนมา{{{{closing_text}}}}")
    doc.add_paragraph()

    # ลายเซ็นผู้ขอ
    doc.add_paragraph(f"{TAB * 6}({{{{requester_name}}}})")
    doc.add_paragraph(f"{TAB * 6}{{{{requester_position}}}}")
    doc.add_paragraph()

    # เรียน + อนุมัติ
    doc.add_paragraph("เรียน ผู้อำนวยการสถาบันพัฒนาและฝึกอบรมโรงงานต้นแบบ")
    doc.add_paragraph(f"{TAB * 2}เพื่อโปรดพิจารณาอนุมัติ")
    doc.add_paragraph()
    doc.add_paragraph(f"ลงชื่อ{DOTS} {TAB}ลงชื่อ{DOTS}")
    doc.add_paragraph(f"({{{{approver_left_name}}}}) {TAB}({{{{approver_right_name}}}})")
    doc.add_paragraph()

    # เอกสารแนบ (sections — engine reuses the shared section parsers)
    doc.add_paragraph(f"ประมาณการค่าใช้จ่าย ({{{{#breakdown}}}})")
    doc.add_paragraph(f"กำหนดการเดินทาง ({{{{#schedule}}}})")

    doc.save(path)
    print("saved:", path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates",
        "บันทึกข้อความ_rsc_blank.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    build(out)

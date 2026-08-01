import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def build_sample_docx(output_path="sample_memo.docx"):
    doc = docx.Document()

    # Title / Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_header.add_run("บันทึกข้อความ")
    run_title.bold = True
    run_title.font.size = Pt(20)

    # Org Info
    p_org = doc.add_paragraph()
    p_org.add_run("ส่วนงาน  ").bold = True
    p_org.add_run("ศูนย์สนับสนุนโครงการหลวงและโครงการพระราชดำริ   โทร 053-218618")

    # Doc No & Date
    p_num_date = doc.add_paragraph()
    p_num_date.add_run("ที่  ").bold = True
    p_num_date.add_run("อว 7608.8.1/1234/69                                    ")
    p_num_date.add_run("วันที่  ").bold = True
    p_num_date.add_run("10 สิงหาคม 2569")

    # Subject
    p_sub = doc.add_paragraph()
    p_sub.add_run("เรื่อง  ").bold = True
    p_sub.add_run("ขออนุมัติเดินทางติดตามงานโครงการเกษตรและระบบควบคุมสภาพแวดล้อม สค.69")

    # Recipient
    p_to = doc.add_paragraph()
    p_to.add_run("เรียน  ").bold = True
    p_to.add_run("ผู้อำนวยการสถาบันพัฒนาและฝึกอบรมโรงงานต้นแบบ")

    # Body Paragraph 1 (Context)
    p_ctx = doc.add_paragraph()
    p_ctx.paragraph_format.first_line_indent = Inches(0.5)
    p_ctx.add_run(
        "ตามที่ ศูนย์สนับสนุนโครงการหลวงและโครงการพระราชดำริ ได้ดำเนินงานโครงการวิจัยและพัฒนาการผลิตสตอเบอรี่และพืชผักอัจฉริยะระบบควบคุมอัจฉริยะ ประจำปีงบประมาณ 2569 นั้น"
    )

    # Body Paragraph 2 (Objective & Location & Schedule)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.first_line_indent = Inches(0.5)
    p_obj.add_run(
        "ในการนี้ คณะทำงานจึงใคร่ขออนุมัติเดินทางไปปฏิบัติงานติดตามงานและทดสอบระบบ ณ ศูนย์พัฒนาโครงการหลวงแม่แฮ จ.เชียงใหม่ ในระหว่างวันที่ 15 สิงหาคม 2569 ถึงวันที่ 20 สิงหาคม 2569 โดยมีรายละเอียดค่าใช้จ่ายดังรายการด้านล่าง"
    )

    # Budget Breakdown Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "รายการ"
    hdr_cells[1].text = "จำนวน / รายละเอียด"
    hdr_cells[2].text = "จำนวนเงิน (บาท)"

    items = [
        ("ค่าเช่าพาหนะ", "จำนวน 6 วัน x 1,500 บาท", "9,000"),
        ("ค่าน้ำมันเชื้อเพลิง", "จำนวน 6 วัน x 1,000 บาท", "6,000"),
        ("ค่าทางด่วน / ผ่านทาง", "จำนวน 1 เหมา", "1,200"),
        ("ค่าเบี้ยเลี้ยงและที่พัก", "เหมาจ่าย", "5,000"),
    ]

    for item, detail, amount in items:
        row_cells = table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = detail
        row_cells[2].text = amount

    # Total Budget Paragraph
    p_budget = doc.add_paragraph()
    p_budget.paragraph_format.first_line_indent = Inches(0.5)
    p_budget.add_run("รวมเป็นเงินทั้งสิ้น  ").bold = True
    p_budget.add_run("21,200 บาท (สองหมื่นเอ็ดพันสองร้อยบาทถ้วน) ")
    p_budget.add_run("เบิกจ่ายจากงบประมาณโครงการฯ")

    # Closing
    p_close = doc.add_paragraph()
    p_close.paragraph_format.first_line_indent = Inches(0.5)
    p_close.add_run("จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ")

    # Signature Block
    p_sig1 = doc.add_paragraph()
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig1.add_run("\n\n(นายรณกร อำพันธ์ศรี)\n")
    p_sig1.add_run("วิศวกร")

    doc.save(output_path)
    print(f"Sample document saved successfully to {output_path}")

if __name__ == "__main__":
    build_sample_docx()
